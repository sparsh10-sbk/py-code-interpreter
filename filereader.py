
from abc import ABC, abstractmethod
import logging
import pdfplumber
import openpyxl
import pandas as pd
from docx import Document 
import pypandoc



logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

class FileReader(ABC):
    @abstractmethod
    def read(self, file_path):
        pass

class PDFReader(FileReader):
    def read(self, file_path):
        try:
            with pdfplumber.open(file_path) as pdf:
                content = ""
                for page in pdf.pages:
                    content += page.extract_text()
            logging.info("PDF file read successfully")
            return content
        except Exception as e:
            logging.error(f"Error reading PDF file: {e}")
            return "Error reading PDF file"

class XLSXReader(FileReader):
    def read(self, file_path):
        try:
            workbook = openpyxl.load_workbook(file_path)
            sheet = workbook.active
            content = ""
            for row in sheet.iter_rows(values_only=True):
                content += " ".join([str(cell) for cell in row]) + "\n"
            logging.info("XLSX file read successfully")
            return content
        except Exception as e:
            logging.error(f"Error reading XLSX file: {e}")
            return "Error reading XLSX file"

class CSVReader(FileReader):
    def read(self, file_path):
        try:
            df = pd.read_csv(file_path)
            logging.info("CSV file read successfully")
            return df.to_string()
        except Exception as e:
            logging.error(f"Error reading CSV file: {e}")
            return "Error reading CSV file"

# class DOCXReader(FileReader):
#     def read(self, file_path):
#         try:
#             doc = Document(file_path)
#             content = ""
#             for paragraph in doc.paragraphs:
#                 content += paragraph.text + "\n"
#             logging.info("DOCX file read successfully")
#             return content
#         except Exception as e:
#             logging.error(f"Error reading DOCX file: {e}")
#             return "Error reading DOCX file"
# either docx or pyandoc can be used;
# class DOCXReader(FileReader):
#     def read(self, file_path):
#         try:
           
#             content = pypandoc.convert_file(file_path, 'plain', format='docx')
#             logging.info("DOCX file read successfully")
#             return content
#         except Exception as e:
#             logging.error(f"Error reading DOCX file: {e}")
#             return "Error reading DOCX file"
class TXTReader(FileReader):
    def read(self, file_path):
        try:
            with open(file_path, 'r') as file:
                content = file.read()
            logging.info("TXT file read successfully")
            return content
        except Exception as e:
            logging.error(f"Error reading TXT file: {e}")
            return "Error reading TXT file"

class FileReaderFactory:
    @staticmethod
    def get_file_reader(file_type):
        if file_type == 'pdf':
            return PDFReader()
        elif file_type == 'xlsx':
            return XLSXReader()
        elif file_type == 'csv':
            return CSVReader()
        elif file_type == 'docx':
            return DOCXReader()
        elif file_type == 'txt':
            return TXTReader()
        else:
            logging.error(f"Unsupported file type: {file_type}")
            raise ValueError("Unsupported file type")
